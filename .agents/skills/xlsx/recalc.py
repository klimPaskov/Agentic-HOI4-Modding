#!/usr/bin/env python3
"""Recalculate an Excel workbook with LibreOffice and report formula errors."""

from __future__ import annotations

import json
import importlib.util
import os
import platform
import shutil
import subprocess
import sys
from pathlib import Path

EXCEL_ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")


def _macro_directory() -> Path:
    system = platform.system()
    if system == "Windows":
        appdata = os.environ.get("APPDATA")
        base = Path(appdata) if appdata else Path.home() / "AppData" / "Roaming"
        return base / "LibreOffice" / "4" / "user" / "basic" / "Standard"
    if system == "Darwin":
        return Path.home() / "Library" / "Application Support" / "LibreOffice" / "4" / "user" / "basic" / "Standard"
    return Path.home() / ".config" / "libreoffice" / "4" / "user" / "basic" / "Standard"


def setup_libreoffice_macro(soffice: str, timeout: int) -> tuple[bool, str | None]:
    """Install the bounded recalculation macro if it is not already present."""
    macro_dir = _macro_directory()
    macro_file = macro_dir / "Module1.xba"
    try:
        if macro_file.exists() and "RecalculateAndSave" in macro_file.read_text(encoding="utf-8"):
            return True, None
        if not macro_dir.exists():
            subprocess.run(
                [soffice, "--headless", "--terminate_after_init"],
                capture_output=True,
                timeout=min(timeout, 10),
                check=False,
            )
            macro_dir.mkdir(parents=True, exist_ok=True)
        macro_file.write_text(
            '''<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>''',
            encoding="utf-8",
        )
        return True, None
    except (OSError, subprocess.SubprocessError) as exc:
        return False, str(exc)


def _scan_workbook(filename: Path) -> dict[str, object]:
    from openpyxl import load_workbook

    details = {error: [] for error in EXCEL_ERRORS}
    total_errors = 0
    values = load_workbook(filename, data_only=True, read_only=True)
    try:
        for worksheet in values.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue
                    for error in EXCEL_ERRORS:
                        if error in cell.value:
                            details[error].append(f"{worksheet.title}!{cell.coordinate}")
                            total_errors += 1
                            break
    finally:
        values.close()

    formula_count = 0
    formulas = load_workbook(filename, data_only=False, read_only=True)
    try:
        for worksheet in formulas.worksheets:
            for row in worksheet.iter_rows():
                formula_count += sum(
                    1 for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")
                )
    finally:
        formulas.close()

    return {
        "status": "success" if total_errors == 0 else "errors_found",
        "total_errors": total_errors,
        "total_formulas": formula_count,
        "error_summary": {
            error: {"count": len(locations), "locations": locations[:20]}
            for error, locations in details.items()
            if locations
        },
    }


def recalc(filename: str | Path, timeout: int = 30) -> dict[str, object]:
    path = Path(filename).expanduser().resolve()
    if not path.is_file():
        return {"error": f"File does not exist: {path}"}
    if timeout <= 0:
        return {"error": "timeout_seconds must be positive"}
    if importlib.util.find_spec("openpyxl") is None:
        return {"error": "Python package openpyxl is required for workbook verification"}

    soffice = shutil.which("soffice")
    if not soffice:
        return {"error": "LibreOffice soffice was not found on PATH"}

    configured, error = setup_libreoffice_macro(soffice, timeout)
    if not configured:
        return {"error": f"Failed to configure LibreOffice macro: {error or 'unknown error'}"}

    command = [
        soffice,
        "--headless",
        "--norestore",
        "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
        str(path),
    ]
    try:
        completed = subprocess.run(command, capture_output=True, text=True, timeout=timeout, check=False)
    except subprocess.TimeoutExpired:
        return {"error": f"LibreOffice recalculation timed out after {timeout} seconds"}
    except OSError as exc:
        return {"error": f"LibreOffice recalculation failed to start: {exc}"}

    if completed.returncode != 0:
        message = completed.stderr.strip() or completed.stdout.strip() or "unknown LibreOffice error"
        return {"error": message, "returncode": completed.returncode}

    try:
        return _scan_workbook(path)
    except ModuleNotFoundError as exc:
        if exc.name == "openpyxl":
            return {"error": "Python package openpyxl is required for workbook verification"}
        return {"error": f"Workbook verification dependency is unavailable: {exc}"}
    except Exception as exc:
        return {"error": f"Workbook verification failed: {exc}"}


def main(argv: list[str] | None = None) -> int:
    args = list(sys.argv[1:] if argv is None else argv)
    if not args or args[0] in {"-h", "--help"}:
        print("Usage: python recalc.py <excel_file> [timeout_seconds]")
        print("Recalculates formulas with LibreOffice and reports Excel errors as JSON.")
        return 0 if args else 1
    try:
        timeout = int(args[1]) if len(args) > 1 else 30
    except ValueError:
        print(json.dumps({"error": "timeout_seconds must be an integer"}, indent=2))
        return 2
    result = recalc(args[0], timeout)
    print(json.dumps(result, indent=2))
    return 1 if "error" in result or result.get("status") == "errors_found" else 0


if __name__ == "__main__":
    raise SystemExit(main())
