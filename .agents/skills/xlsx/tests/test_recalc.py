from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook


MODULE_PATH = Path(__file__).resolve().parents[1] / "recalc.py"
SPEC = importlib.util.spec_from_file_location("hoi4_xlsx_recalc", MODULE_PATH)
assert SPEC and SPEC.loader
RECALC = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(RECALC)


class RecalcTests(unittest.TestCase):
    def test_missing_file_is_reported(self) -> None:
        result = RECALC.recalc("definitely-missing-workbook.xlsx")
        self.assertIn("does not exist", result["error"])

    def test_missing_soffice_is_reported(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "book.xlsx"
            Workbook().save(workbook_path)
            with mock.patch.object(RECALC.shutil, "which", return_value=None):
                result = RECALC.recalc(workbook_path)
        self.assertEqual(result["error"], "LibreOffice soffice was not found on PATH")

    def test_missing_openpyxl_is_reported_before_libreoffice_use(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "book.xlsx"
            Workbook().save(workbook_path)
            with mock.patch.object(RECALC.importlib.util, "find_spec", return_value=None):
                with mock.patch.object(RECALC.shutil, "which") as which:
                    result = RECALC.recalc(workbook_path)
        self.assertEqual(result["error"], "Python package openpyxl is required for workbook verification")
        which.assert_not_called()

    def test_scan_counts_formulas_without_errors(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "book.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 2
            sheet["A2"] = "=A1*2"
            workbook.save(workbook_path)
            result = RECALC._scan_workbook(workbook_path)
        self.assertEqual(result["status"], "success")
        self.assertEqual(result["total_formulas"], 1)
        self.assertEqual(result["total_errors"], 0)

    def test_nonpositive_timeout_is_rejected_before_provider_use(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "book.xlsx"
            Workbook().save(workbook_path)
            result = RECALC.recalc(workbook_path, timeout=0)
        self.assertEqual(result["error"], "timeout_seconds must be positive")


if __name__ == "__main__":
    unittest.main()
